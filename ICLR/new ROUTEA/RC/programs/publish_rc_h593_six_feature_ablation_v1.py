#!/usr/bin/env python3
"""Build a new complete-report edition only after the five-fold ablation validates."""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'reports/new_hyp_complete_results_20260916_v1'
OUT = ROOT / 'reports/new_hyp_complete_results_20260920_ablation_v1'
RESULT = ROOT / 'results/rc_h593_six_feature_ablation_v1'
ARCHIVE = 'new_HYP_complete_results_20260920_ablation.zip'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n')


def validate(directory):
    manifest=json.loads((directory/'archive_manifest.json').read_text())
    for name, meta in manifest['files'].items():
        p=directory/name
        assert p.stat().st_size==meta['bytes'] and sha(p)==meta['sha256'], name
    with zipfile.ZipFile(directory/ARCHIVE) as z:
        assert z.testzip() is None
        assert set(z.namelist())==set(manifest['files'])|{'archive_manifest.json'}
        for name in z.namelist():
            assert hashlib.sha256(z.read(name)).hexdigest()==sha(directory/name), name
    for name in ['all_results.xlsx','complete_results_zh.docx','H593_ablation_zh.docx']:
        with zipfile.ZipFile(directory/name) as z:
            assert z.testzip() is None
    book=openpyxl.load_workbook(directory/'all_results.xlsx',read_only=True)
    assert book['H593_ablation'].max_row==28
    assert book['H593_ablation_groups'].max_row==25
    book.close()


def main():
    result=json.loads((RESULT/'result.json').read_text())
    validation=json.loads((RESULT/'validation.json').read_text())
    assert result['status']=='H593_SIX_FEATURE_ABLATION_COMPLETE'
    assert validation['status']=='H593_ABLATION_INDEPENDENT_ACTION_COUNTS_PASS'
    assert validation['result']['sha256']==sha(RESULT/'result.json')
    assert (RESULT/'report_zh.md').is_file()
    if OUT.exists():
        validate(OUT)
        assert sha(OUT/'H593_ablation_evidence/result.json')==sha(RESULT/'result.json')
        print('EXISTING_COMPLETE_ABLATION_PACKAGE_VERIFIED');return
    with tempfile.TemporaryDirectory(prefix='.h593_ablation_publish_',dir=ROOT/'reports') as temp:
        stage=Path(temp)
        for p in SOURCE.iterdir():
            if p.is_file() and p.suffix=='.zip':continue
            if p.is_dir():shutil.copytree(p,stage/p.name)
            else:shutil.copy2(p,stage/p.name)
        previous={str(p.relative_to(SOURCE)):sha(p) for p in SOURCE.rglob('*') if p.is_file() and p.suffix!='.zip'}
        shutil.copy2(stage/'archive_manifest.json',stage/'historical_archive_manifest_20260919.json')
        shutil.copy2(stage/'validation.json',stage/'historical_validation_before_ablation_20260920.json')
        evidence=stage/'H593_ablation_evidence'
        shutil.copytree(RESULT,evidence)
        report=(RESULT/'report_zh.md').read_text()
        (stage/'H593_ablation_zh.md').write_text(report+'\n[原始证据目录](H593_ablation_evidence/)\n')
        complete=stage/'complete_results_zh.md'
        complete.write_text(complete.read_text()+'\n\n## 2026-09-20 补充：H593 六项逐删与重训\n\n'+
                            '\n'.join(report.splitlines()[2:])+'\n\n[逐折原始分数及验证](H593_ablation_evidence/)\n')
        readme=stage/'README_zh.md'
        readme.write_text('# new HYP 结果汇总包：2026-09-20 消融补充版\n\n'+
                          '[H593六项消融](H593_ablation_zh.md) · [Excel](all_results.xlsx) · [原始证据](H593_ablation_evidence/)\n\n'+
                          '本版保留2026-09-19汇总包的全部非ZIP内容，并增加H593的60个删项重训头及冻结置零对照；旧统计和图片不改写。\n\n'+
                          '\n'.join(readme.read_text().splitlines()[1:]))
        book=openpyxl.load_workbook(stage/'all_results.xlsx')
        sheet=book.create_sheet('H593_ablation')
        columns=['model','correct','rescue_vs_RAW','break_vs_RAW','rescue_vs_FULL','break_vs_FULL','changed_vs_FULL','switches','MRR']
        sheet.append(columns)
        for name,s in result['summary'].items():sheet.append([name]+[s[c] for c in columns[1:]])
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
        sheet=book.create_sheet('H593_ablation_groups')
        sheet.append(['model','full_minus_ablation_equal_component','bootstrap95_lower','bootstrap95_upper','one_sided_p','Holm_p_six'])
        for name,s in result['summary'].items():
            if 'bootstrap95' in s:sheet.append([name,s['equal_component_full_minus_ablation'],*s['bootstrap95'],s['signflip_one_sided_p'],s['holm_adjusted_p_within_six']])
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
        book.save(stage/'all_results.xlsx')
        for base in ['complete_results_zh','H593_ablation_zh']:
            for ext in ['html','docx']:
                args=['pandoc',base+'.md','--standalone','--toc','--toc-depth=2','-o',base+'.'+ext]
                if ext=='html':args+=['--metadata','title=new HYP 实验汇总']
                subprocess.run(args,cwd=stage,check=True)
        scope=json.loads((stage/'package_scope.json').read_text())
        scope['updated']='2026-09-20'
        scope['included'].append('H593 OOF5 six-coordinate frozen zeroing and leave-one-out COST1/CE refits; all fold scores and validations')
        scope['ablation_primary']='COST1';scope['ablation_is_external_confirmation']=False
        scope['ablation_optimizer_updates_including_full_and_fresh_replays']=280000
        write(stage/'package_scope.json',scope)
        sources=json.loads((stage/'source_manifest.json').read_text())
        for p in RESULT.rglob('*'):
            if p.is_file():sources[str(p.relative_to(ROOT))]=dict(sha256=sha(p),bytes=p.stat().st_size)
        write(stage/'source_manifest.json',sources)
        checks=json.loads((stage/'validation.json').read_text())
        checks.update(updated='2026-09-20',status='COMPLETE_RESULTS_PLUS_H593_ABLATION_PACKAGED',
                      ablation_validation=validation, historical_source_directory=str(SOURCE),
                      no_new_figures=True, no_previous_predictions_changed=True,
                      archive_validation='archive_manifest.json')
        write(stage/'validation.json',checks)
        files={str(p.relative_to(stage)):dict(bytes=p.stat().st_size,sha256=sha(p)) for p in sorted(stage.rglob('*'))
               if p.is_file() and p.name!='archive_manifest.json'}
        manifest=dict(status='COMPLETE_RESULTS_WITH_H593_ABLATION_ARCHIVE',updated='2026-09-20',
                      previous_source=str(SOURCE),previous_source_hashes=previous,
                      new_source_result_sha256=sha(RESULT/'result.json'),files=files)
        write(stage/'archive_manifest.json',manifest)
        with zipfile.ZipFile(stage/ARCHIVE,'w',zipfile.ZIP_DEFLATED) as z:
            for name in sorted(set(files)|{'archive_manifest.json'}):z.write(stage/name,name)
        validate(stage)
        shutil.copytree(stage,OUT)
    validate(OUT)
    print(json.dumps(dict(status='COMPLETE_ABLATION_PACKAGE_VERIFIED',directory=str(OUT),archive=str(OUT/ARCHIVE),sha256=sha(OUT/ARCHIVE)),ensure_ascii=False))


if __name__=='__main__':main()
