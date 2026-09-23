#!/usr/bin/env python3
"""Publish measured head timing and the user-approved RPC evidence reclassification."""
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/new_hyp_complete_results_20260916_v1'
BENCH = ROOT / 'results/rc_head_training_time_v1_20260918'
BACKUP = ROOT / 'reports/new_hyp_complete_results_20260916_v1_backups/before_training_time_20260918/new_HYP_complete_results_20260916.zip'
STAGE = ROOT / 'reports/new_hyp_complete_results_training_time_build_20260918'
NEW_SOURCES = {}
TABLES = {}


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def read(p):
    p = Path(p)
    NEW_SOURCES[str(p.relative_to(ROOT))] = dict(sha256=sha(p), bytes=p.stat().st_size)
    return json.loads(p.read_text())


def write(p, obj):
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False)+'\n')


def csvout(name, rows):
    TABLES[name] = rows
    with (STAGE / (name+'.csv')).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)


def table(header, rows):
    return '\n'.join(['| '+' | '.join(header)+' |', '| '+' | '.join(['---']*len(header))+' |'] +
                     ['| '+' | '.join(map(str,r))+' |' for r in rows])


def fit(x, y):
    x, y = np.asarray(x, dtype=float), np.asarray(y, dtype=float)
    slope, intercept = np.polyfit(x, y, 1)
    prediction = intercept+slope*x
    ss = float(((y-y.mean())**2).sum())
    return dict(intercept_seconds=float(intercept), slope_seconds_per_effective_query=float(slope),
                R_squared=1-float(((prediction-y)**2).sum())/ss if ss else None,
                RMSE_seconds=float(np.sqrt(((prediction-y)**2).mean())),
                measured_x=x.tolist(), median_seconds=y.tolist(), residual_seconds=(y-prediction).tolist())


def measured_tables(d):
    csvout('training_time_repeats', d['records'])
    summary=[]; fits=[]
    for model in ('COST1','CE'):
        mm=[]
        for n in (32,128,593):
            rr=[r for r in d['records'] if r['model']==model and r['pool_size']==n]
            assert len(rr)==3 and len({r['theta_sha256'] for r in rr})==1
            ts=np.array([r['fit_seconds'] for r in rr])
            m=dict(model=model,pool_size=n,effective_queries=rr[0]['effective_queries'],
                   C128_absent=rr[0]['absent_from_C128'],repeats=3,steps=2000,
                   median_seconds=float(np.median(ts)),minimum_seconds=float(ts.min()),maximum_seconds=float(ts.max()),
                   mean_seconds=float(ts.mean()),sample_sd_seconds=float(ts.std(ddof=1)),
                   timing_scope='Warm CPU original train call only; cache and startup excluded')
            summary.append(m);mm.append(m)
        f=fit([m['effective_queries'] for m in mm],[m['median_seconds'] for m in mm])
        fits.append(dict(model=model,**{k:v for k,v in f.items() if not isinstance(v,list)},
                         effective_queries=json.dumps(f['measured_x']),
                         median_seconds=json.dumps(f['median_seconds']),residual_seconds=json.dumps(f['residual_seconds']),
                         scope='Three observed scales only; no extrapolation or universal complexity claim'))
    csvout('training_time_summary',summary);csvout('training_time_fits',fits)
    log=ROOT/'logs/rc_original7_train128_readout-5139291.out'
    NEW_SOURCES[str(log.relative_to(ROOT))]=dict(sha256=sha(log),bytes=log.stat().st_size)
    history=[];counts={}
    for line in log.read_text().splitlines():
        if not line.startswith('{'):continue
        row=json.loads(line)
        if row.get('event')!='FIXED_RECIPE_HEAD_FITTED':continue
        model=row['model'];counts[model]=counts.get(model,0)+1
        history.append(dict(model=model,FULL_queries=row['effective_FULL_rows'],PAIR_auxiliary_rows=64,
            call=counts[model],fit_seconds=row['seconds'],parameter_sha256=row['parameter_sha256'],
            job_id='5139291',node='hkn0002',CPU_threads=8,steps=2000,
            source='logs/rc_original7_train128_readout-5139291.out',
            scope='Historical mixed PAIR64 plus FULL recipe; not comparable to current COST1/CE curve'))
    assert len(history)==6
    csvout('historical_training_time',history)
    sched=read(BENCH/'historical_scheduler.json')['records']
    csvout('historical_job_elapsed',sched)
    return summary,fits,history


def plotting(summary, fits, history):
    os.environ.setdefault('MPLCONFIGDIR','/tmp/rc-head-time-mpl')
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    font_manager.fontManager.addfont('/usr/share/fonts/google-droid-sans-fonts/DroidSansFallbackFull.ttf')
    plt.rcParams.update({'font.family':['DejaVu Sans','Droid Sans Fallback'],'axes.unicode_minus':False,
                         'pdf.fonttype':42,'svg.fonttype':'path','font.size':11})
    fig,axes=plt.subplots(1,2,figsize=(16,6.5),facecolor='#fafcfe')
    for ax in axes:
        ax.set_facecolor('#fafcfe');ax.spines[['top','right']].set_visible(False)
        ax.grid(axis='y',alpha=.16);ax.set_axisbelow(True);ax.set_ylabel('训练函数耗时（秒）')
    ax=axes[0]
    for model,color in [('COST1','#087f8c'),('CE','#c18b32')]:
        rr=[r for r in summary if r['model']==model];f=next(f for f in fits if f['model']==model)
        x=np.array([r['effective_queries'] for r in rr]);y=np.array([r['median_seconds'] for r in rr])
        err=np.array([[r['median_seconds']-r['minimum_seconds'] for r in rr],
                      [r['maximum_seconds']-r['median_seconds'] for r in rr]])
        xx=np.linspace(x.min(),x.max(),250)
        ax.plot(xx,f['intercept_seconds']+f['slope_seconds_per_effective_query']*xx,color=color,linestyle='--',alpha=.8)
        ax.errorbar(x,y,yerr=err,fmt='o',color=color,capsize=5,markersize=7,label=model+'：中位数及范围')
        for r in rr:
            reps=[v['fit_seconds'] for v in TABLES['training_time_repeats'] if v['model']==model and v['pool_size']==r['pool_size']]
            ax.scatter([r['effective_queries']]*len(reps),reps,color=color,s=13,alpha=.65)
        for xi,yi in zip(x,y):ax.annotate(f'{yi:.3f}s',(xi,yi),xytext=(5,11 if model=='COST1' else -18),textcoords='offset points',color=color,fontsize=10)
    ax.set_xticks([31,120,570],['32\n有效31','128\n有效120','593\n有效570'])
    ax.set_xlim(0,640);ax.set_ylim(0,max(r['maximum_seconds'] for r in summary)*1.22)
    ax.set_title('当前同配方计时：32 / 128 / 593 选入图',fontsize=15,pad=15)
    ax.set_xlabel('选入图数；拟合 N 使用实际进入 loss 的图数')
    ax.legend(loc='upper left',fontsize=10,frameon=False)
    ax=axes[1];xx=np.array([32,96,128]);yy=np.array([np.median([r['fit_seconds'] for r in history if r['FULL_queries']==n]) for n in xx])
    f=fit(xx,yy);line=np.linspace(32,128,100)
    ax.plot(line,f['intercept_seconds']+f['slope_seconds_per_effective_query']*line,'--',color='#596d86',label='历史同配方线性拟合')
    for n,y in zip(xx,yy):
        values=[r['fit_seconds'] for r in history if r['FULL_queries']==n]
        ax.scatter([n]*len(values),values,color='#596d86',s=30)
        ax.annotate(f'{y:.2f}s',(n,y),xytext=(0,12),textcoords='offset points',ha='center')
    ax.set_xticks(xx);ax.set_xlim(20,142);ax.set_ylim(0,max(r['fit_seconds'] for r in history)*1.22)
    ax.set_title('历史计时：FULL32 / 96 / 128 + PAIR64',fontsize=15,pad=15)
    ax.set_xlabel('FULL 图数（每次另有固定 PAIR64 辅助训练）');ax.legend(frameon=False,loc='upper left',fontsize=10)
    fig.suptitle('七参数头的训练时间：分清同配方趋势与历史成本',fontsize=20,y=.97)
    fig.text(.05,.055,'左：同一节点、8 CPU 线程、FP64、AdamW 2000步，每配置重复3次；缓存和首次预热不计入。',color='#526577',fontsize=11)
    fig.text(.05,.02,'右：历史训练与独立重放各1次。两侧配方和实现不同，不跨面板连线；均不能代表编码 + RoMa 的端到端耗时。',color='#526577',fontsize=11)
    fig.tight_layout(rect=[.02,.105,.99,.92],w_pad=4)
    for ext in ['png','pdf','svg']:fig.savefig(STAGE/('training_time_scaling.'+ext),dpi=190,facecolor=fig.get_facecolor())
    plt.close(fig)
    fig,axes=plt.subplots(1,3,figsize=(17,5.8),facecolor='#fafcfe')
    panels=[('H593：开发 OOF5',593,[426,440,481,486]),('GroZi480：正式外部确认',480,[321,326,355,367]),('ISIC537：跨域探索',537,[466,469,500,513])]
    for ax,(name,n,values) in zip(axes,panels):
        bars=ax.bar(['RAW','COST4','COST1','CE'],np.array(values)/n*100,color=['#a8b5c5','#507dad','#087f8c','#c18b32'],width=.65)
        ax.set_facecolor('#fafcfe');ax.set_title(name,pad=17,fontsize=14);ax.set_ylim(0,109);ax.set_ylabel('准确率 (%)')
        ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',alpha=.13);ax.set_axisbelow(True)
        for b,v in zip(bars,values):ax.text(b.get_x()+b.get_width()/2,b.get_height()+1.2,f'{v}/{n}\n{v/n:.2%}',ha='center',fontsize=10)
    fig.suptitle('主要结果：内部开发、正式外部确认与跨域探索分别呈现',fontsize=19,y=.98)
    fig.text(.045,.02,'RPC 依据 2026-09-18 的参考图适用性决定移入诊断附录，不纳入正式外部证据。各面板分母与头来源独立。',fontsize=10,color='#526577')
    fig.tight_layout(rect=[.02,.065,.99,.92],w_pad=2)
    for ext in ['png','pdf','svg']:fig.savefig(STAGE/('main_results.'+ext),dpi=180,facecolor=fig.get_facecolor())
    plt.close(fig)


def timing_text(d,summary,fits,history):
    lines=['## 3.1 新增：32／128／593 的训练耗时（2026-09-18）','',
        '本节回答训练小头要花多少时间。采用原 H593 执行顺序的嵌套前32／128／593张，在同一节点调用原 COST1/CE 训练函数；两模型均为6权重＋1偏置、零初始化、FP64、AdamW 2000步、8 CPU线程。此处32/128是H593内部的计时子集，不是重跑历史 EVAL32/EVAL128，也不是新的准确率学习曲线。','',
        '自然C128中缺target的1／8／23张沿用原loss排除规则，所以实际训练数是31／120／570。全H593重放的参数与原冻结COST1/CE逐位一致；计时参数不保存为新模型。','',
        table(['模型','选入图','有效训练图','3次中位数/s','最小—最大/s'],[[r['model'],r['pool_size'],r['effective_queries'],f"{r['median_seconds']:.3f}",f"{r['minimum_seconds']:.3f}—{r['maximum_seconds']:.3f}"] for r in summary]),'',
        '![训练耗时趋势：当前同配方与历史混合配方分别拟合](training_time_scaling.png)','',
        '拟合使用有效训练图数 N，T 的单位为秒；每种模型仅用三个规模的中位数进行仿射最小二乘拟合。误差棒是3次实测的最小—最大范围，不是置信区间。','']
    for f in fits:
        lines.append(f"- {f['model']}：T(N) = {f['intercept_seconds']:.6f} + {f['slope_seconds_per_effective_query']:.6f}N；R²={f['R_squared']:.5f}，RMSE={f['RMSE_seconds']:.5f}s。")
    lines += ['', '只有三个观测规模，拟合仅描述本节点、本实现、固定2000步的观测范围，不证明普适复杂度、不向更大样本量外推。','',
        f"计时环境：{d['runtime']['hostname']}，{d['runtime']['cpu_model']}；PyTorch {d['runtime']['torch']}，CPU 8线程。申请GPU是为了运行短时分区，本次训练未用GPU。",'',
        f"纯头计时从已有内存特征开始，包含优化器创建和2000次更新；不包含排队、Python/库启动、特征缓存读取、编码器/RoMa前向和评价。缓存读取与组批另耗 {read(BENCH/'data_closure.json')['cache_load_and_batch_seconds']:.3f}s；首个 COST1 热身调用 {d['warmup'][0]['seconds']:.3f}s、CE 热身 {d['warmup'][1]['seconds']:.3f}s，均未纳入趋势拟合。不要把预热后的约3秒误写成完整RAW＋RoMa训练流水线总耗时。",'',
        '### 历史实际耗时，独立保留','',
        table(['历史配方','FULL图','额外PAIR','首次训练/s','独立重放/s'],[[m,n,64,*[f"{r['fit_seconds']:.3f}" for r in history if r['model']==m]] for m,n in [('ORIGINAL7',32),('FULL96',96),('FULL128',128)]]),'',
        '历史记录来自 job5139291、CPU节点hkn0002、8线程；每次仍为2000步，但包含固定PAIR64辅助项及不同训练实现。其整个作业为299秒，包含三套头及其独立重放、读取和评价，不能分配成一个头的训练耗时。历史H593单模型作业COST1为60秒、CE为59秒，包含加载、两次拟合和验证；原日志未单列纯train耗时，因此不与历史32/128拼成一条曲线。','',
        '[逐次实测](training_time_repeats.csv) · [汇总](training_time_summary.csv) · [拟合系数/残差](training_time_fits.csv) · [历史纯训练](historical_training_time.csv) · [历史作业总耗时](historical_job_elapsed.csv)','']
    return '\n'.join(lines)


def manual_section():
    path=ROOT/'results/rc_crisp_manual_baseline_v1/result.json';d=read(path)
    v=read(path.parent/'result_validation.json');assert v['result']['sha256']==sha(path)
    h=d['panels']['H593'];out=[]
    for m,c in h['counts'].items():
        rr=h['comparisons']['RAW'].get(m,dict(rescues=0,breaks=0,net=0))
        out.append(dict(panel='H593',model=m,queries=593,correct=c,accuracy_percent=c/593*100,
                        rescues=rr['rescues'],breaks=rr['breaks'],net=rr['net'],evidence='Grouped OOF development; trained heads versus fixed heuristics'))
    csvout('manual_head_H593',out)
    lines=['## 3.2 训练与手填参数：审稿问题的直接对照','',
        '固定同一六维输入、RAW自然C128和正logit才SWITCH的规则，手填EQUAL=[1,1,1,1,1,1,0]、BALANCED=[1,.2,.2,.2,.2,.2,0]；数值在读本轮结果前冻结。全零头严格回到RAW，仅作实现检查。H593训练头使用原分组OOF结果，没有用全H593拟合准确率冒充留出。','',
        table(['方法','正确/593','对RAW救回','改错','净增'],[[r['model'],r['correct'],r['rescues'],r['breaks'],r['net']] for r in out]),'',
        'COST1相对两种手工规则均净增22/593；按64个来源连通组重采样的图像微平均差95%区间：相对EQUAL为[0.162,7.290]pp，相对BALANCED为[1.563,5.942]pp。CE相对两种规则均净增27/593。以上是所测规则下的开发证据，不证明所有免训练规则都不可能达到同样效果，也没有新的GroZi手填头对照。','',
        'CRISP采用作者公开评分函数，适配image-to-image、同一自然C128、image-only tokens；它在H593为373/593，匹配token口径的普通MaxSim为424/593。这是当前适配的结果，不是原文text-to-flowchart任务的完全复现。RPC相关读出仅见诊断附录，不用于上述训练价值论证。','',
        '训练与手填在部署时执行相同的线性头，训练增加的是一次参数估计成本。当前证据支持“小成本学习在H593上优于所测固定规则”，不能泛化成“任何任务都必须训练”。','']
    return '\n'.join(lines),d


def reclassify_rpc(text, manual):
    old = re.search(r'### RPC600\n(.*?)\n### ISIC537', text, flags=re.S)
    assert old
    rpc_models = old.group(1)
    text = text[:old.start()] + text[old.end()-len('### ISIC537'):]
    rpc_comparisons = [s for s in text.splitlines() if s.startswith('| RPC600 |') and '→' in s]
    rpc_inventory = [s for s in text.splitlines() if s.startswith('| RPC / COST1 |')]
    rpc_actions = [s for s in text.splitlines() if s.startswith('| RPC | 51 |')]
    # Replace the combined headline matrix with formal and exploratory panels only.
    start=text.index('## 10. 外部全部模型与控制结果')
    stop=text.index('### GroZi480',start)
    models=('RAW','COST4','GROUP_COST4','COST1','CE','RAW2_CE','COST1_CBIND','CE_CBIND','COST4_CBIND','GROUP_COST4_CBIND','RAW2_CE_CBIND')
    grozi=read(ROOT/'results/rc_new_hyp_grozi120_external_v1/result.json')
    isic=read(ROOT/'results/rc_new_hyp_isic_transfer_v1/result.json')
    new='## 10. 正式外部确认与跨域探索的全部模型\n\n'+table(
        ['模型 / 对照','GroZi /480：正式外部','ISIC /537：探索'],
        [[m,grozi['counts'][m],isic['counts'][m]] for m in models])+'\n\n'
    text=text[:start]+new+text[stop:]
    text='\n'.join(s for s in text.splitlines() if not re.match(r'^\| RPC(?:600| / COST1| \|)',s))+'\n'
    changes={
      '截至 2026-09-16，Europe/Berlin。':'原结果整理截至 2026-09-16；训练耗时与证据范围更新于 2026-09-18，Europe/Berlin。',
      'new HYP 的技术改进已经在 H593 分组开发和两个外部商品数据集上出现。':'new HYP 的技术改进已在 H593 分组开发和 GroZi 商品外部确认中出现；ISIC作为跨域探索，RPC仅保留参考图不足诊断。',
      '本次整理不训练、不推理、不换头、不重设门。':'原2026-09-16整理未训练或推理。本次仅新增同配方耗时重放、已完成的固定手工规则对照与RPC证据范围修订；不替换现有头，不改写原准确率。',
      '外部三集共享同一批全 H593 训练并冻结的头；':'三个已运行数据库面板共享同一批全 H593 训练并冻结的头，其中RPC不计入正式证据；',
      '随后同一固定主头在GroZi/RPC确认，ISIC探索也正向':'随后同一固定主头在GroZi确认，ISIC探索也正向；RPC已移入诊断',
      '## 9. 外部数据库：数据构建、图库和证据范围':'## 9. 正式外部与探索数据库：数据构建、图库和证据范围',
      '三套均使用同一批全H593固定头':'GroZi、ISIC及诊断RPC均使用同一批全H593固定头',
      'RPC当前协议是单商品跨相机，不是多商品结账检测；reference视角/物体像素可能影响难度，但没有通过新增reference受控实验量化，不能断言低分全由reference不好造成。':'RPC原协议只使用camera0图作为参考侧，现按参考图适用性决定移入诊断附录，不作为正式外部确认。',
      '主模型相对RAW：GroZi34救0损、净+34（+7.08pp）；RPC18救3损、净+15（+2.50pp）；ISIC34救0损、净+34（+6.33pp）。RPC的3次损失完整计入，沿用可靠净增目标；另两套观察到零损失不构成普遍零损保证。':'主模型相对RAW：正式GroZi为34救0损、净+34（+7.08pp）；探索ISIC为34救0损、净+34（+6.33pp）。两套观察到零损失不构成普遍零损保证。RPC的完整救损保留在诊断附录。',
      'GroZi以27来源视频等权、RPC以200 SKU在17大类内分层、ISIC以346患者等权；':'GroZi以27来源视频等权、ISIC以346患者等权；',
      'COST1分别可靠优于RAW、同规模COST4、GROUP_COST4、RAW-only两参数对照以及自己的CBIND。CE在GroZi/RPC比COST1多12/10张，但两处组区间跨0，继续是预定次模型。':'正式GroZi中，COST1分别可靠优于RAW、同规模COST4、GROUP_COST4、RAW-only两参数对照以及自己的CBIND。CE在GroZi比COST1多12张，但组区间跨0，继续是预定次模型。',
      'RPC按相机：RAW→COST1→CE为camera1 66→75→79 /200；camera2 31→35→39 /200；camera3 95→97→99 /200。三个点估计都有净增，但没有据此宣称每相机显著；camera2明显更困难。':'RPC的相机分组读出已移入诊断附录。',
      'RPC只有34.50%主模型准确率，GO指可靠净增与绑定门通过，不能写成已经具备普遍高精度。公开数据未知预训练暴露未排除；ISIC近重复/拍摄会话独立性未由唯一图片哈希保证。':'公开数据未知预训练暴露未排除；ISIC近重复/拍摄会话独立性未由唯一图片哈希保证。RPC历史GO不再作为当前正式证据。',
      '在同表示/同小头结构下产生内部和跨数据集净增。':'在同表示/同小头结构下产生H593内部和GroZi正式外部净增，ISIC另作探索。',
      '[RPC原结果]':'[RPC历史诊断原结果]',
      '外部11臂、配对区间':'正式外部/探索11臂、独立RPC诊断、配对区间',
    }
    for old,new in changes.items():
        assert old in text, old
        text=text.replace(old,new)
    scope='**2026-09-18 范围修订：RPC缺少本任务认可的信息充分的专用参考图，退出正式外部确认与训练必要性论证；保留为诊断附录。GroZi是正式商品外部确认，ISIC仍是已打开队列上的跨域探索。**\n\n'
    text=text.replace('## 1. 所有数字先对齐模型与数据',scope+'## 1. 所有数字先对齐模型与数据')
    rpc=manual['panels']['RPC600'];newrows=[]
    for m,c in rpc['counts'].items():
        x=rpc['comparisons']['RAW'].get(m,dict(rescues=0,breaks=0,net=0))
        newrows.append([m,c,x['rescues'],x['breaks'],x['net']])
    appendix=['## 附录D：RPC参考图不足诊断（不计入正式实验）','',
        '范围决定日期：2026-09-18，依据用户对参考图适用性的要求。原协议在操作上确实有200张camera0参考侧图片，但这种跨相机参考侧不被认可为当前任务所需的信息充分的专用reference。原600图读出全部保留；历史“外部GO”标签属于当时协议记录，当前不沿用。','',
        '这项范围调整在结果已打开后记录，不删除低分或损失，也不将RPC纳入正式跨数据集成功或训练必要性统计。未进行补充reference的受控比较，故不能断言所有错误均由reference不足造成。','',
        '### 历史冻结头的完整读出','',rpc_models.strip(),'',
        '### 历史配对区间（诊断）','',
        table(['面板','基线→模型','救/损','净增','等SKU差 pp','95%组区间 pp'],
              [[c.strip() for c in s.strip('|').split('|')] for s in rpc_comparisons]),'',
        '以上历史区间按200 SKU、17大类内分层重采样，仅为既有诊断统计。','',
        '### 新增免训练评分读出（诊断）','',
        table(['方法','正确/600','对RAW救回','改错','净增'],newrows),'',
        '手填平衡规则209/600、COST1为207/600，配对净差的SKU分组95%区间为[-2.00,2.67]pp；保留这个读出，不用于选择论文主模型。当前没有证明训练COST1在此诊断面板优于手填平衡规则。','',
        '历史错误与动作：COST1共393错，25例target缺C128，368例候选内仍错；80例最高challenger为target但HOLD，288例为其余候选内错误。51次SWITCH中18救回、3改错、30错换另一个错；549次HOLD。','',
        '相机分组RAW→COST1→CE：camera1为66→75→79 /200；camera2为31→35→39 /200；camera3为95→97→99 /200。','',
        '[历史模型CSV](RPC_diagnostic_models.csv) · [历史比较CSV](RPC_diagnostic_comparisons.csv) · [新增免训练CSV](RPC_manual_diagnostic.csv) · [范围决定](rpc_evidence_scope.json)','']
    text+='\n'+'\n'.join(appendix)
    return text


def workbooks_and_scope_csvs():
    book=openpyxl.load_workbook(STAGE/'all_results.xlsx')
    for name,newname in [('external_all_models','RPC_diagnostic_models'),('external_comparisons','RPC_diagnostic_comparisons')]:
        p=STAGE/(name+'.csv')
        with p.open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
        rpc=[r for r in rows if r['panel']=='RPC600'];rest=[r for r in rows if r['panel']!='RPC600']
        assert rpc
        for r in rpc:r['evidence_level']='Reference-insufficiency diagnostic only; excluded from formal external evidence on 2026-09-18'
        csvout(name,rest);csvout(newname,rpc)
    for name,rows in TABLES.items():
        if name in book:del book[name]
        sheet=book.create_sheet(name[:31]);sheet.append(list(rows[0]))
        for row in rows:sheet.append([v if not isinstance(v,(list,dict)) else json.dumps(v,ensure_ascii=False) for v in row.values()])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in list(sheet.columns):sheet.column_dimensions[col[0].column_letter].width=24
    book.save(STAGE/'all_results.xlsx')
    return book.sheetnames


def main():
    assert not (OUT/'training_time_update_validation.json').exists(), 'Update already published'
    v=read(BENCH/'validation.json');d=read(BENCH/'result.json')
    assert v['result']['sha256']==sha(BENCH/'result.json') and v['timed_fits']==18
    assert d['status']=='HEAD_TRAINING_TIMING_COMPLETE' and all(
        r['historical_head_max_abs_error']==0 for r in d['records'] if r['pool_size']==593)
    STAGE.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(BACKUP) as z:
        assert z.testzip() is None
        for name in z.namelist():
            assert '/' not in name and '\\' not in name
            (STAGE/name).write_bytes(z.read(name))
    oldvalidation=json.loads((STAGE/'validation.json').read_text())
    write(STAGE/'historical_validation_20260916.json',oldvalidation)
    summary,fits,history=measured_tables(d)
    plotting(summary,fits,history)
    manualtext,manual=manual_section()
    rpc_rows=[]
    for m,c in manual['panels']['RPC600']['counts'].items():
        p=manual['panels']['RPC600']['comparisons']['RAW'].get(m,dict(rescues=0,breaks=0,net=0))
        rpc_rows.append(dict(panel='RPC600_DIAGNOSTIC',model=m,correct=c,queries=600,
                            rescues=p['rescues'],breaks=p['breaks'],net=p['net'],formal_external_evidence=False))
    csvout('RPC_manual_diagnostic',rpc_rows)
    timing=timing_text(d,summary,fits,history)
    text=(STAGE/'complete_results_zh.md').read_text()
    text=reclassify_rpc(text,manual)
    text=text.replace('## 4. 旧模型与固定 EVAL32/128：完整保留收益和取舍',timing+'\n'+manualtext+'\n## 4. 旧模型与固定 EVAL32/128：完整保留收益和取舍')
    text=text.replace('- [核心结果条形图PDF]', '- [训练耗时曲线PDF](training_time_scaling.pdf)、[SVG](training_time_scaling.svg)、[PNG](training_time_scaling.png)，以及[计时说明](training_time_zh.md)。\n- [核心结果条形图PDF]')
    (STAGE/'complete_results_zh.md').write_text(text)
    (STAGE/'training_time_zh.md').write_text('# 七参数训练耗时与手工规则比较\n\n'+timing+'\n'+manualtext)
    scope=read(ROOT/'registry/rc_rpc_evidence_scope_update_20260918.json')
    write(STAGE/'rpc_evidence_scope.json',scope)
    write(STAGE/'training_time_source.json',d)
    write(STAGE/'training_time_fit.json',dict(fits=fits,scope='Observed warmed CPU head timing only',extrapolated=False))
    sheets=workbooks_and_scope_csvs()
    for ext in ['html','docx']:
        args=['pandoc','complete_results_zh.md','--standalone','--toc','--toc-depth=2']
        if ext=='html':args+=['--metadata','title=药盒与外部结果：训练耗时与证据范围更新','--include-in-header=style.html']
        subprocess.run(args+['-o','complete_results_zh.'+ext],cwd=STAGE,check=True)
    for p in [Path(__file__).resolve(),ROOT/'programs/benchmark_rc_head_training_time_v1.py',
              ROOT/'registry/rc_head_training_time_authority_v1_20260918.json',
              ROOT/'plan/RC_HEAD_TRAINING_TIME_SCALING_V1_20260918.md']:
        NEW_SOURCES[str(p.relative_to(ROOT))]=dict(sha256=sha(p),bytes=p.stat().st_size)
    manifest=json.loads((STAGE/'source_manifest.json').read_text());manifest.update(NEW_SOURCES)
    write(STAGE/'source_manifest.json',manifest)
    newvalidation=dict(oldvalidation,updated='2026-09-18',
        status='RESULTS_DOSSIER_WITH_MEASURED_TRAINING_TIME_AND_RPC_SCOPE_UPDATE',
        original_count_verification='Preserved 2026-09-16 checks; source result metrics unchanged',
        timing_validation=v, timing_training_updates=d['timed_optimizer_updates']+d['warmup_optimizer_updates'],
        timing_is_accuracy_experiment=False,RPC_formal_external_confirmation=False,
        scope_decision='rpc_evidence_scope.json',source_count=len(manifest),
        bootstrap_intervals='Sealed result intervals reused; no new accuracy resampling')
    write(STAGE/'validation.json',newvalidation)
    for name in ['all_results.xlsx','complete_results_zh.docx']:
        with zipfile.ZipFile(STAGE/name) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if n.endswith('.xml'):ET.fromstring(z.read(n))
    book=openpyxl.load_workbook(STAGE/'all_results.xlsx',read_only=True)
    for name in ['external_all_models','external_comparisons']:
        assert all(row[0]!='RPC600' for row in book[name].values)
    assert len(list(book['training_time_repeats'].values))==19
    assert len(list(book['training_time_summary'].values))==7
    assert 'RPC_diagnostic_models' in book
    from PIL import Image
    for name in ['main_results.png','training_time_scaling.png']:
        with Image.open(STAGE/name) as im:im.verify()
    preappendix=text.split('## 附录D：')[0]
    assert '| RPC600 |' not in preappendix and '| RPC |' not in preappendix and '两个外部商品数据集' not in text
    assert 'training_time_scaling.png' in (STAGE/'complete_results_zh.html').read_text()
    with zipfile.ZipFile(STAGE/'complete_results_zh.docx') as z:
        assert len([n for n in z.namelist() if n.startswith('word/media/')])>=2
        assert '32／128／593' in z.read('word/document.xml').decode()
    update=dict(status='MEASURED_TIMING_DOCUMENTS_AND_RPC_SCOPE_PASS',
        backup_archive=dict(path=str(BACKUP),sha256=sha(BACKUP)),timing_fits=fits,
        measured_fits=18,repeats_per_configuration=3,full_head_parameter_max_abs_error=0,
        main_plot_excludes_RPC=True,formal_tables_exclude_RPC=True,RPC_diagnostics_retained=True,
        docx_xlsx_xml_valid=True,spreadsheet_sheets=sheets,
        source_sha256={str(p.relative_to(ROOT)):sha(p) for p in [BENCH/'result.json',BENCH/'validation.json']},
        artifact_sha256={p.name:sha(p) for p in sorted(STAGE.iterdir()) if p.is_file()})
    write(STAGE/'training_time_update_validation.json',update)
    archive=STAGE/'new_HYP_complete_results_20260916.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(STAGE.iterdir()):
            if p.is_file() and p!=archive:z.write(p,p.name)
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
        for name in z.namelist():assert hashlib.sha256(z.read(name)).hexdigest()==sha(STAGE/name)
    # Publish the reviewed documents first and atomically replace the requested ZIP last.
    for p in STAGE.iterdir():
        if p.is_file() and p!=archive:
            tmp=OUT/('.'+p.name+'.publishing');shutil.copy2(p,tmp);os.replace(tmp,OUT/p.name)
    tmp=OUT/'.new_HYP_complete_results_20260916.zip.publishing'
    shutil.copy2(archive,tmp);os.replace(tmp,OUT/archive.name)
    print(json.dumps(dict(status='UPDATED',archive=str(OUT/archive.name),archive_sha256=sha(OUT/archive.name),
                          summary=summary,fits=fits),ensure_ascii=False))


if __name__=='__main__':main()
