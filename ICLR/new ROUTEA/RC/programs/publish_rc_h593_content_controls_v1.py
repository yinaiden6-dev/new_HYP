#!/usr/bin/env python3
"""Append validated learned content controls to the ablation report edition."""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

import openpyxl

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'reports/new_hyp_complete_results_20260920_ablation_v1'
OUT=ROOT/'reports/new_hyp_complete_results_20260920_controls_v1'
RESULT=ROOT/'results/rc_h593_learned_colnomic_only_v1'
ARCHIVE='new_HYP_complete_results_20260920_controls.zip'
AUTH=ROOT/'registry/rc_h593_content_controls_package_authority_v1_20260920.json'


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):return json.loads(p.read_text())
def write(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2,allow_nan=False)+'\n')


def validate(directory):
    manifest=read(directory/'archive_manifest.json')
    for name,b in manifest['files'].items():
        p=directory/name;assert p.stat().st_size==b['bytes'] and sha(p)==b['sha256'],name
    with zipfile.ZipFile(directory/ARCHIVE) as z:
        assert z.testzip() is None and set(z.namelist())==set(manifest['files'])|{'archive_manifest.json'}
        for name in z.namelist():assert hashlib.sha256(z.read(name)).hexdigest()==sha(directory/name),name
    for name in ('complete_results_zh.docx','H593_content_controls_zh.docx','all_results.xlsx'):
        with zipfile.ZipFile(directory/name) as z:assert z.testzip() is None
    book=openpyxl.load_workbook(directory/'all_results.xlsx',read_only=True)
    assert book['H593_content_controls'].max_row==10 and book['H593_content_groups'].max_row==5
    assert 'H593_ablation' in book.sheetnames
    book.close()


def main():
    for source in read(AUTH)['sources'].values():
        assert sha(Path(source['path']))==source['sha256'],source['path']
    result=read(RESULT/'result.json');v=read(RESULT/'validation.json')
    assert result['status']=='H593_LEARNED_COLNOMIC_ONLY_COMPLETE'
    assert v['status']=='CONTENT_INDEPENDENT_ACTION_RANK_COUNTS_PASS' and v['result']['sha256']==sha(RESULT/'result.json')
    # Verify the complete preceding edition, not just its existence.
    import publish_rc_h593_six_feature_ablation_v1 as previous
    previous.validate(SOURCE)
    if OUT.exists():
        validate(OUT);assert sha(OUT/'H593_content_evidence/result.json')==sha(RESULT/'result.json');print('EXISTING_CONTENT_PACKAGE_VERIFIED');return
    with tempfile.TemporaryDirectory(prefix='.h593_content_package_',dir=ROOT/'reports') as temp:
        stage=Path(temp)
        for p in SOURCE.iterdir():
            if p.is_file() and p.suffix=='.zip':continue
            if p.is_dir():shutil.copytree(p,stage/p.name)
            else:shutil.copy2(p,stage/p.name)
        shutil.copy2(stage/'archive_manifest.json',stage/'historical_archive_manifest_before_content_20260920.json')
        shutil.copy2(stage/'validation.json',stage/'historical_validation_before_content_20260920.json')
        shutil.copytree(RESULT,stage/'H593_content_evidence')
        report=(RESULT/'report_zh.md').read_text()
        (stage/'H593_content_controls_zh.md').write_text(report+'\n[逐候选数据与验证](H593_content_evidence/)\n')
        complete=stage/'complete_results_zh.md'
        complete.write_text(complete.read_text()+'\n\n## 2026-09-20 补充：learned ColNomic-only correction\n\n'+'\n'.join(report.splitlines()[2:])+'\n')
        readme=stage/'README_zh.md'
        readme.write_text('# new HYP 完整汇总：六项消融与纯内容学习对照\n\n[纯内容对照](H593_content_controls_zh.md) · [六项消融](H593_ablation_zh.md) · [全部表格](all_results.xlsx)\n\n'+
                         '\n'.join(readme.read_text().splitlines()[1:]))
        book=openpyxl.load_workbook(stage/'all_results.xlsx');sheet=book.create_sheet('H593_content_controls')
        sheet.append(['model','correct','MRR','rescue_vs_RAW','break_vs_RAW','changed_vs_RAW'])
        for m,s in result['summary'].items():sheet.append([m,s['correct'],s['MRR'],s['against_RAW']['rescue'],s['against_RAW']['loss'],s['against_RAW']['changed']])
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
        sheet=book.create_sheet('H593_content_groups');sheet.append(['comparison','full_rescue','full_break','equal_component_full_minus_content','bootstrap95_lower','bootstrap95_upper'])
        for m,s in result['comparisons'].items():sheet.append([m,s['rescue'],s['loss'],s['equal_component_difference'],*s['bootstrap95']])
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;book.save(stage/'all_results.xlsx')
        for base in ('complete_results_zh','H593_content_controls_zh'):
            for ext in ('html','docx'):
                args=['pandoc',base+'.md','--standalone','--toc','--toc-depth=2','-o',base+'.'+ext]
                if ext=='html':args+=['--metadata','title=new HYP 实验汇总']
                subprocess.run(args,cwd=stage,check=True)
        scope=read(stage/'package_scope.json');scope['included'].append('H593 learned ColNomic-only CONTENT7/MAXSIM3 correction, COST1/CE, 75 feature shards and OOF5 validated predictions')
        scope['content_controls_optimizer_updates_with_replays']=100000;scope['content_controls_external_confirmation']=False;write(stage/'package_scope.json',scope)
        sources=read(stage/'source_manifest.json')
        for p in RESULT.rglob('*'):
            if p.is_file():sources[str(p.relative_to(ROOT))]=dict(sha256=sha(p),bytes=p.stat().st_size)
        write(stage/'source_manifest.json',sources)
        validation=read(stage/'validation.json');validation.update(status='COMPLETE_RESULTS_WITH_ABLATION_AND_CONTENT_CONTROLS',content_controls_validation=v,previous_edition=str(SOURCE));write(stage/'validation.json',validation)
        files={str(p.relative_to(stage)):dict(bytes=p.stat().st_size,sha256=sha(p)) for p in sorted(stage.rglob('*')) if p.is_file() and p.name!='archive_manifest.json'}
        write(stage/'archive_manifest.json',dict(status='COMPLETE_ABLATION_CONTENT_ARCHIVE',previous_edition=str(SOURCE),content_result_sha256=sha(RESULT/'result.json'),files=files))
        with zipfile.ZipFile(stage/ARCHIVE,'w',zipfile.ZIP_DEFLATED) as z:
            for name in sorted(set(files)|{'archive_manifest.json'}):z.write(stage/name,name)
        validate(stage);shutil.copytree(stage,OUT)
    validate(OUT);print(json.dumps(dict(status='CONTENT_CONTROLS_PACKAGE_VERIFIED',archive=str(OUT/ARCHIVE),sha256=sha(OUT/ARCHIVE)),ensure_ascii=False))


if __name__=='__main__':main()
