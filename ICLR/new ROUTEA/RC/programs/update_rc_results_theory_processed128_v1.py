#!/usr/bin/env python3
"""Append published theory sources and the completed processed128 regression."""
import argparse
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET

os.environ.setdefault('MPLCONFIGDIR', '/tmp/rc-theory-processed128-mpl')
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/new_hyp_complete_results_20260916_v1'
ARCHIVE = 'new_HYP_complete_results_20260916.zip'
BASE_SHA = '0035b0a5a131531b40cf4df50e44d8decaaa31996a5df700cfeab34f9406f401'
BACKUP = ROOT / 'reports/new_hyp_complete_results_20260916_v1_backups/before_theory_processed128_20260919' / ARCHIVE
PROC = ROOT / 'results/rc_new_hyp_processed128_regression_v1'
RECEIPTS = ROOT / 'results/rc_results_theory_processed128_publication_v1_20260919'
THEORY = [
    'reports/NEW_HYP_THEORY_DEFINITION_AND_PROPOSITIONS_V1_20260909.md',
    'reports/NEW_HYP_THEORY_DEFINITION_AND_PROPOSITIONS_V2_20260910.md',
    'reports/NEW_HYP_UNIFIED_DECISION_HYPOTHESIS_V3_20260911.md',
    'reports/NEW_HYP_PRODUCT_CONTRAST_MATHEMATICAL_REVIEW_V1_20260909.md',
    'reports/REPORT_NEW_HYP_CALIBRATION_MECHANISM_SYNTHESIS_V1_20260909.md',
    'reports/NEW_HYP_PRIMARY_LITERATURE_SCOPE_V1_20260909.md',
    'plan/RC_NEW_HYP_PAPER_SCOPE_FREEZE_V1_20260915.md',
]
SOURCES = {}
TABLES = {}


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def write(p, d):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(d, ensure_ascii=False, indent=2, allow_nan=False)+'\n')


def bind(p):
    p = Path(p)
    SOURCES[str(p.relative_to(ROOT))] = dict(sha256=sha(p), bytes=p.stat().st_size)
    return SOURCES[str(p.relative_to(ROOT))]


def read(p):
    bind(p)
    return json.loads(Path(p).read_text())


def copy_source(p, dest):
    bind(p)
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(p, dest)
    assert sha(p) == sha(dest)


def csvout(stage, name, rows):
    TABLES[name] = rows
    with (stage / (name+'.csv')).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)


def table(headers, rows):
    return '\n'.join(['| '+' | '.join(headers)+' |', '| '+' | '.join(['---']*len(headers))+' |']+
                     ['| '+' | '.join(map(str, r))+' |' for r in rows])


def processed(stage):
    d = read(PROC / 'result.json'); v = read(PROC / 'result_validation.json')
    assert d['status'] == 'PROCESSED128_FROZEN_REGRESSION_COMPLETE'
    assert v['status'] == 'PROCESSED128_SEALS_COUNTS_ACTIONS_PASS'
    assert v['result']['sha256'] == sha(PROC / 'result.json')
    assert len(d['rows']) == d['queries'] == 128 and d['training_updates'] == 0
    assert not d['formal_GO_claimed'] and d['training_image_byte_overlap'] == 0
    assert len({r['identity'] for r in d['rows']}) == 128
    assert sum(r['target_in_C128'] for r in d['rows']) == d['target_recall_C128'] == 127
    assert sum(r['training_identity_overlap'] for r in d['rows']) == 5
    for b in v['predictions']:
        p = Path(b['path']); assert sha(p) == b['sha256']
        copy_source(p, stage / 'processed128_evidence/prediction_validations' / (p.parent.name+'.json'))
    assert len(v['predictions']) == 16
    for name in ['result.json', 'result_validation.json', 'per_query.csv', 'report_zh.md',
                 'selection.json', 'metadata_validation.json', 'all_predictions_prejoin_seal.json']:
        copy_source(PROC / name, stage / 'processed128_evidence' / name)
    a = Path(v['authority']['path']); assert sha(a) == v['authority']['sha256']
    copy_source(a, stage / 'processed128_evidence/authority.json')
    order = ['RAW', 'COST4', 'GROUP_COST4', 'COST1', 'CE', 'RAW2_CE']
    order += sorted(set(d['counts'])-set(order))
    all_models = []
    for m in order:
        n = sum(r['correct'][m] for r in d['rows'])
        rescue = sum(not r['correct']['RAW'] and r['correct'][m] for r in d['rows'])
        loss = sum(r['correct']['RAW'] and not r['correct'][m] for r in d['rows'])
        assert (n,rescue,loss) == (d['counts'][m],d['rescues'][m],d['breaks'][m])
        switches = sum(not r['holds'][m] for r in d['rows'])
        all_models.append(dict(panel='processed128 synthetic regression', model=m, queries=128,
            correct=n, accuracy_percent=n/128*100, rescue_vs_RAW=rescue, loss_vs_RAW=loss,
            net_vs_RAW=rescue-loss, switches=switches, holds=128-switches,
            target_recall_C128=127, head_lineage=d['head_lineage'], candidate_source=d['candidate_source'],
            action=d['action'], evidence_level='Outcome-blind synthetic regression; not independent external confirmation'))
    csvout(stage, 'processed128_all_models', all_models)
    comparisons=[]
    for m,c in d['comparisons'].items():
        comparisons.append(dict(panel='processed128 synthetic regression', model=m, baseline='RAW',
            rescues=c['rescues'], breaks=c['breaks'], net=c['net'],
            accuracy_difference_pp=c['accuracy_difference']*100,
            group_CI95_low_pp=c['component_bootstrap95'][0]*100,
            group_CI95_high_pp=c['component_bootstrap95'][1]*100,
            groups=c['components'], reliable_net_positive=c['reliable_net_positive']))
    csvout(stage, 'processed128_comparisons', comparisons)
    for key, name, field in [('by_corruption','processed128_by_corruption','corruption'),
                             ('by_training_identity_overlap','processed128_by_overlap','training_identity_overlap')]:
        rows=[]
        for group, stats in d[key].items():
            subset=[r for r in d['rows'] if (str(r[field]) if field=='training_identity_overlap' else r[field])==group]
            assert len(subset)==stats['queries']
            for m in order:
                n=sum(r['correct'][m] for r in subset)
                rescue=sum(not r['correct']['RAW'] and r['correct'][m] for r in subset)
                loss=sum(r['correct']['RAW'] and not r['correct'][m] for r in subset)
                assert (n,rescue,loss)==(stats['counts'][m],stats['rescues'][m],stats['breaks'][m])
                rows.append({field:group, 'model':m, 'queries':len(subset), 'correct':n,
                             'rescues':rescue, 'breaks':loss, 'target_recall_C128':stats['target_recall_C128']})
        csvout(stage,name,rows)
    changed=[]
    for m in ['COST1','CE']:
        for r in d['rows']:
            if r['selected'][m] == r['selected']['RAW']: continue
            changed.append(dict(model=m,query_id=r['query_id'],origin_name=r['origin_name'],
                original_path=r['original_path'],corruption=r['corruption'],
                training_identity_overlap=r['training_identity_overlap'],target_in_C128=r['target_in_C128'],
                RAW_correct=r['correct']['RAW'],model_correct=r['correct'][m],
                change='rescue' if r['correct'][m] else ('break' if r['correct']['RAW'] else 'wrong_to_wrong'),
                RAW_selected=r['selected']['RAW'],model_selected=r['selected'][m]))
    csvout(stage,'processed128_changed_queries',changed)
    with (PROC/'per_query.csv').open(encoding='utf-8-sig') as f: rows=list(csv.DictReader(f))
    assert len(rows)==128
    csvout(stage,'processed128_per_query',rows)
    for m in ['COST1','CE']:
        assert not d['comparisons'][m]['reliable_net_positive']
    return d,v,all_models,comparisons


def plot_processed(stage, d):
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt, font_manager
    font_manager.fontManager.addfont('/usr/share/fonts/google-droid-sans-fonts/DroidSansFallbackFull.ttf')
    plt.rcParams.update({'font.family':['DejaVu Sans','Droid Sans Fallback'], 'axes.unicode_minus':False,
                         'pdf.fonttype':42, 'svg.fonttype':'path', 'font.size':11})
    fig,axes=plt.subplots(1,2,figsize=(14,5.8),facecolor='#fafcfe',gridspec_kw={'width_ratios':[1.5,1]})
    models=['RAW','COST4','COST1','CE'];colors=['#a8b5c5','#507dad','#087f8c','#c18b32']
    ax=axes[0];values=[d['counts'][m] for m in models]
    bars=ax.bar(models,[v/128*100 for v in values],color=colors,width=.6)
    for b,v in zip(bars,values):ax.text(b.get_x()+b.get_width()/2,b.get_height()+1.8,f'{v}/128\n{v/128:.2%}',ha='center')
    ax.set_ylim(0,112);ax.set_yticks([0,20,40,60,80,100]);ax.set_ylabel('准确率 (%)')
    ax.set_title('原 121 个正确全部保留',pad=18,fontsize=15)
    ax=axes[1];x=[0,1,2]
    for i,m in enumerate(models[1:]):
        rescue=d['rescues'][m];loss=d['breaks'][m]
        ax.bar(i-.16,rescue,width=.3,color='#087f8c',label='救回' if i==0 else None)
        ax.bar(i+.16,loss,width=.3,color='#b54d55',label='改错' if i==0 else None)
        for xx,yy in [(i-.16,rescue),(i+.16,loss)]:ax.text(xx,yy+.07,str(yy),ha='center')
    ax.set_xticks(x,models[1:]);ax.set_ylim(0,3);ax.set_yticks([0,1,2,3]);ax.set_ylabel('图片数')
    ax.set_title('相对 RAW：小幅净增，未观察到损失',pad=18,fontsize=14);ax.legend(frameon=False,loc='upper left')
    for ax in axes:
        ax.set_facecolor('#fafcfe');ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',alpha=.15);ax.set_axisbelow(True)
    fig.suptitle('processed128：冻结 full-H593 头的合成处理图回归',fontsize=19,y=.98)
    fig.text(.04,.055,'同一自然 C128（target 召回127/128）；128图／128身份；无重训。主模型 COST1，次模型 CE。',color='#526577')
    fig.text(.04,.015,'COST1／CE 净增95%区间均包含0；不是独立外部确认，也不构成普遍零损失保证。',color='#526577')
    fig.tight_layout(rect=[.015,.11,.99,.91],w_pad=3)
    for ext in ['png','pdf','svg']:fig.savefig(stage/('processed128_results.'+ext),dpi=180,facecolor=fig.get_facecolor())
    plt.close(fig)


def processed_text(d, all_models, comparisons):
    lines=['## 3.3 processed128：高 RAW 准确率下的冻结头回归（2026-09-19归档）','',
        '从 `ap7811-benchmark/1/processed` 在查看结果前按固定hash选入128张合成处理图，对应128个不同reference身份。使用2026-09-13冻结的full-H593头，零新增训练、零阈值调整。全库5413物理reference／5412校正身份，RAW自然C128后评价全部127个challenger；最高logit大于0才SWITCH，否则保留RAW。该面板独立列账，不与旧EVAL128的88→99或H593五折合并。','',
        table(['模型','正确/128','准确率','对RAW救回／改错','净增','SWITCH'],
              [[r['model'],r['correct'],f"{r['accuracy_percent']:.2f}%",f"{r['rescue_vs_RAW']}／{r['loss_vs_RAW']}",f"{r['net_vs_RAW']:+d}",r['switches']] for r in all_models if not r['model'].endswith('_CBIND')]),'',
        '![processed128准确率与救回、改错](processed128_results.png)','',
        '**RAW原正确121张全部保留：COST1救回1张，CE救回2张。** COST1共2次SWITCH（1救回、1错换另一个错），CE共3次SWITCH（2救回、1错换另一个错）。COST4、GROUP_COST4和RAW2_CE均没有正确数净增。','',
        '自然C128包含target为127/128，1张正确候选缺席并保留在分母中。COST1剩余6错＝1张候选缺席＋5张候选内判断错；CE剩余5错＝1＋4。','',
        table(['相对RAW的比较','净差/128','准确率差 pp','来源组bootstrap 95%区间 pp'],
              [[r['model'],r['net'],f"{r['accuracy_difference_pp']:+.4f}",f"[{r['group_CI95_low_pp']:.4f}, {r['group_CI95_high_pp']:.4f}]"] for r in comparisons if r['model'] in ('COST1','CE')]),'',
        '两组区间下界均为0，所以只报告观察到小幅净增和零改错，不宣称可靠正净增或总体永不损失。此面板是合成处理图回归，不是独立外部确认。与H593训练图的字节重叠为0，但有5个身份重叠；其余123身份上RAW116、COST1 117、CE118，新增救回全部来自该非重叠部分。字节不重复不等于来源或语义完全独立。','',
        '### 实际救回与候选绑定对照','',
        '- `PROC-Q-0099`：`50mLcarton__motion_blur+aged_05.jpg`，COST1和CE均救回。',
        '- `PROC-Q-0091`：`cla04-0002-05__aged+motion_blur_02.jpg`，CE额外救回。','',
        table(['模型','正确绑定/128','候选证据错绑/128'],
              [[m,d['counts'][m],d['counts'][m+'_CBIND']] for m in ['COST4','GROUP_COST4','COST1','CE','RAW2_CE']]),'',
        '上述错绑是候选级联合证据控制，不能作为空间ownership或每个token精确对应必要性的证明。按处理类型及训练身份重叠分组的所有臂保留在CSV/Excel，未按最好子集筛选。','',
        '[全部11臂](processed128_all_models.csv) · [比较与区间](processed128_comparisons.csv) · [逐图结果](processed128_per_query.csv) · [所有COST1/CE改变决策](processed128_changed_queries.csv) · [处理类型](processed128_by_corruption.csv) · [训练身份重叠](processed128_by_overlap.csv)','',
        '[原始结果JSON](processed128_evidence/result.json) · [原验证记录](processed128_evidence/result_validation.json) · [完整原报告](processed128_evidence/report_zh.md)。本次归档复算128图、11臂的正确数、救损和分组计数，并核对16片验证SHA；未重训、未重新推理。','']
    return '\n'.join(lines)


def theory(stage):
    entries=[]
    for source in THEORY:
        p=ROOT/source;dest=stage/'theory_sources'/p.name
        copy_source(p,dest)
        entries.append(dict(title=p.read_text().splitlines()[0].lstrip('# '), source=source,
            packaged_path='theory_sources/'+p.name, sha256=sha(p), bytes=p.stat().st_size,
            scope='Historical original, byte-preserved; current claims follow the 2026-09-19 summary and RPC scope update'))
    csvout(stage,'theory_source_index',entries)
    v2=(ROOT/THEORY[1]).read_text()
    core=v2[v2.index('## 1. 定义与实际实现'):v2.index('## 7. 扩大样本证据')]
    v3=(ROOT/THEORY[2]).read_text()
    decision=v3[v3.index('## 统一对象与任务'):v3.index('## 同一解释对应已有正、负结果')]
    intro='''# new HYP：理论定义、命题与当前证据说明

归档日期：2026-09-19。本文件把理论定义、五条形式化性质、决策目标及当前证据放在一起，完整历史原文另附于 `theory_sources/`，其SHA与仓库原件一致。

**当前名称为 new HYP — Reference-conditioned Joint Evidence Calibration for Fine-grained Retrieval。** 研究对象是候选reference条件下的身份判断；旧空间/superregion Reference HYP的失败保留，ownership不属于当前已经完成的主张。实现仍须披露RoMa、ColNomic和共享七参数头，不能把抽象命名等同于跨编码器证明。

任务训练只用query/reference身份或正负检索关系，冻结基础模型的预训练另行披露。五条性质是带条件的结构、信息或代数结论，不是普遍识别保证，也不把基础代数或共享参数可扩展性宣称为首次发现。

以下“定义与性质”完整摘入2026-09-10 V2第1–6节；其中实验语句保留其历史时点。随后列出已完成的新证据，避免把旧文中的“当前”“尚未完成”当成2026-09-19状态。

'''
    current='''## 当前证据与理论的对应（2026-09-19）

| 理论层次 | 已完成的对应证据 | 边界 |
|---|---|---|
| 共享reference条件证据与校准 | 同一full-H593固定COST1头，GroZi480为RAW321→355，34救0损；ISIC537为466→500，34救0损 | GroZi为正式外部确认；ISIC为已打开队列的探索。不是任意新身份保证 |
| 信息接口是否保留原评分 | 缺失完整reference质量统计的接口不能普遍重建；补齐后可无损恢复原评分 | 恢复原评分不自动提升识别，也不证明旧空间P成立 |
| 动作目标与错误绝对压负的差异 | 同H593五折、同七参数，COST4 440→COST1 481，44救3损；CE486 | 训练成本单因素有净增，CE相对COST1优势的组区间跨0；不宣称普适最优成本 |
| 实际参数学习的价值 | 同H593自然C128，手填等权／平衡均459，COST1 481、CE486 | 对所测两种固定规则成立；COST4仅440，不能说一切训练版都优于免训练 |
| 高基线保持回归 | processed128：RAW121、COST1 122、CE123；保住原121正确 | 合成处理回归；净增区间包含0，不是外部GO或普遍零损失定理 |
| 候选绑定与空间所有权 | 已有整包候选证据错绑对照；processed128中COST1 122→92、CE123→74 | 支持候选级联合绑定；不证明token对应、mask或ownership |

旧EVAL32的25→28和旧EVAL128的88→99来自固定ORIGINAL7；difficult90的61→69来自另一FROZEN_C头。H593结果是五折OOF，外部和processed128使用冻结full-H593头。各面板和头不混成一条样本量学习曲线。

RPC按2026-09-18决定仅保留参考图不足诊断。附带2026-09-15旧范围文档中的RPC“外部确认”是历史表述，已被当前范围修订覆盖。完整过目不忘系统、未知拒识、持续注册后的旧身份保持、多物体ownership及新编码器训练仍是未来方向。

## 完整原文入口

以下文件按原字节保存。V1/V2/V3记录不同阶段；V3是工作假说及可证伪预测，不因收入本包就升级成普适理论定理。原文内部指向其它仓库证据的链接仍可能需要原工作区。

'''
    index=table(['原文','日期版本与用途','包内文件'],[[e['title'],Path(e['source']).stem,'['+Path(e['source']).name+']('+e['packaged_path']+')'] for e in entries])
    text=intro+core+'\n'+decision+'\n'+current+index+'\n\n[返回当前主总结](complete_results_zh.md) · [原文SHA索引](theory_source_index.csv)\n'
    (stage/'theory_zh.md').write_text(text)
    (stage/'theory_sources/README_zh.md').write_text('# 历史理论原文\n\n本目录7份原文保持仓库字节不变；历史状态和RPC旧范围不可替代当前结论。先阅读[当前理论说明](../theory_zh.md)和[主总结](../complete_results_zh.md)。\n')
    return entries


def render(stage, basename):
    for ext in ['html','docx']:
        args=['pandoc',basename+'.md','--standalone','--toc','--toc-depth=2']
        if ext=='html':args+=['--metadata','title=new HYP：完整理论与processed128补充','--include-in-header=style.html']
        subprocess.run(args+['-o',basename+'.'+ext],cwd=stage,check=True)


def verify(stage):
    manifest=json.loads((stage/'archive_manifest.json').read_text())
    expected=set(manifest['files'])|{'archive_manifest.json'}
    actual={p.relative_to(stage).as_posix() for p in stage.rglob('*') if p.is_file() and p.name!=ARCHIVE}
    assert expected==actual,(expected-actual,actual-expected)
    for name,info in manifest['files'].items():assert sha(stage/name)==info['sha256'],name
    for source,info in manifest['new_sources'].items():assert sha(ROOT/source)==info['sha256'],source
    for entry in manifest['theory_sources']:assert sha(stage/entry['packaged_path'])==entry['sha256']
    for filename in ['complete_results_zh.docx','theory_zh.docx','processed128_zh.docx','all_results.xlsx']:
        with zipfile.ZipFile(stage/filename) as z:
            assert z.testzip() is None
            for name in z.namelist():
                if name.endswith('.xml'):ET.fromstring(z.read(name))
    book=openpyxl.load_workbook(stage/'all_results.xlsx',read_only=True)
    assert book['processed128_all_models'].max_row==12
    assert book['processed128_per_query'].max_row==129
    assert book['processed128_by_overlap'].max_row==23
    assert book['theory_source_index'].max_row==8
    for name in ['external_all_models','external_comparisons']:
        assert all(row[0]!='RPC600' for row in book[name].values)
    text=(stage/'complete_results_zh.md').read_text()
    assert '## 3.3 processed128' in text and '## 12.1 理论定义' in text
    html_tables=re.findall(r'<table\b[^>]*>.*?</table>',(stage/'complete_results_zh.html').read_text(),flags=re.S)
    assert 'processed128' in html_tables[0], 'processed128 must be a row in the headline table'
    assert '| RPC600 |' not in text.split('## 附录D：')[0]
    for t in ['MANUAL_EQUAL','CRISP','COST4：T(N)','## 5. 早期空间','## 6. new HYP']:
        assert t in text,t
    with zipfile.ZipFile(stage/'complete_results_zh.docx') as z:
        xml=z.read('word/document.xml').decode()
        assert all(s in xml for s in ['processed128','理论定义','CRISP','1.564'])
        assert len([n for n in z.namelist() if n.startswith('word/media/')])>=3
    from PIL import Image
    with Image.open(stage/'processed128_results.png') as im:im.verify()
    for target in re.findall(r'\]\(([^\n]+?)\)', text):
        target=target.strip('<>')
        if target.startswith(('theory_sources/','processed128_','theory_zh.')):
            assert (stage/target.split('#')[0]).exists(),target
    return manifest


def build():
    assert sha(OUT/ARCHIVE)==BASE_SHA,'Source ZIP changed; review before building'
    BACKUP.parent.mkdir(parents=True,exist_ok=True)
    if not BACKUP.exists():shutil.copy2(OUT/ARCHIVE,BACKUP)
    assert sha(BACKUP)==BASE_SHA
    stage=Path(tempfile.mkdtemp(prefix='new_hyp_theory_processed128_build_20260919_',dir=ROOT/'reports'))
    with zipfile.ZipFile(BACKUP) as z:
        assert z.testzip() is None
        for name in z.namelist():
            assert len(Path(name).parts)==1
            (stage/name).write_bytes(z.read(name))
    # Keep historical whole-package hash snapshots separate from the current one.
    for name in ['validation.json','training_time_update_validation.json','training_time_cost4_update_validation.json']:
        dest=stage/'historical_validation'/name;dest.parent.mkdir(exist_ok=True);shutil.copy2(stage/name,dest)
    original_manifest=json.loads((stage/'source_manifest.json').read_text())
    d,v,models,comparisons=processed(stage)
    plot_processed(stage,d)
    ptext=processed_text(d,models,comparisons)
    (stage/'processed128_zh.md').write_text('# processed128：完整冻结头回归读出\n\n'+ptext)
    theory_entries=theory(stage)
    main=(stage/'complete_results_zh.md').read_text()
    main=main.replace('训练耗时与证据范围更新于 2026-09-18','训练耗时与证据范围更新于 2026-09-18；理论原文与processed128归档补齐于 2026-09-19')
    old='原2026-09-16整理未训练或推理。本次仅新增同配方耗时重放、已完成的固定手工规则对照与RPC证据范围修订；不替换现有头，不改写原准确率。'
    assert old in main
    main=main.replace(old,'原2026-09-16整理未训练或推理；2026-09-18加入同配方耗时重放、固定手工规则对照与RPC范围修订。2026-09-19补齐理论原文及已完成的processed128结果，仅复算计数和整理文件，不进行新训练或推理，不替换现有头。')
    old='本次逐图复算覆盖 36 个结果/面板检查条目及原128干预，导出 418 行模型汇总（包含不同账本的重复基线，不是独立实验数）；索引 80 份历史报告。'
    assert old in main
    main=main.replace(old,'2026-09-16底账覆盖36个结果/面板检查条目及原128干预，导出418行模型汇总（含不同账本的重复基线，不是独立实验数）。后续CRISP/手填对照与耗时独立列账；本次另复算processed128的128图、11臂与分组统计，补入完成报告，并保留80份历史报告索引。')
    row='| processed128 合成处理回归 | 121/128 | 同一full-H593固定 COST4 121/128 | COST1 122/128；CE 123/128 | 冻结头无重训；保留RAW原121正确；净增区间包含0，非独立外部确认 |\n'
    marker='\n\n这里的“原配方”均包含'
    assert main.count(marker)==1
    main=main.replace(marker,'\n'+row+'\n这里的“原配方”均包含')
    marker='## 4. 旧模型与固定 EVAL32/128：完整保留收益和取舍'
    main=main.replace(marker,ptext+'\n'+marker)
    theory_section='''## 12.1 理论定义、命题原文与当前解释：已随包补齐

[理论完整阅读版](theory_zh.md)／[HTML](theory_zh.html)／[Word](theory_zh.docx)包含定义、五条性质、决策目标及当前证据对应。完整V1、V2、V3、乘积对比推导、机制收口、文献边界与历史写作范围均保存在 `theory_sources/`，详见[7份原文及SHA索引](theory_source_index.csv)。原文按历史字节保存；当前结果和RPC范围以本总结为准。

| 性质或定义 | 已有形式化内容 | 不能由此推出 |
|---|---|---|
| reference登记与共享参数分离 | 加入reference不增加identity专属参数 | 任意新身份都能被召回、认对 |
| 必要统计量与接口压缩 | 同压缩状态对应不同原评分时，不可普遍重建 | 所有压缩必然降低准确率 |
| 对称相对比较 | 无epsilon时尺度不变；保留epsilon时有明确差异公式 | 实际绝对尺度永远无用 |
| HOLD/SWITCH的正确条件 | target为challenger时需胜HOLD及最强wrong | 必须把每个wrong都压到0以下 |
| 乘积对比的表达范围 | 正值、S=ML、epsilon=0且无floor时，dS提供非仿射基函数；实际保护项另行说明 | 增加这一列必然改善泛化 |

这些结构和数学性质与H593成本对照、候选绑定、GroZi外部确认、ISIC探索以及processed128回归分别对应，不把经验增益升级为普遍识别定理。Ownership和完整过目不忘系统仍保留为未来工作。

'''
    main=main.replace('## 13. 展示与可复核材料',theory_section+'## 13. 展示与可复核材料')
    for e in theory_entries:
        if e['source'].startswith('reports/'):
            main=main.replace('../'+Path(e['source']).name,e['packaged_path'])
        else:main=main.replace('../../'+e['source'],e['packaged_path'])
    marker='- [完整统计Excel](all_results.xlsx)'
    main=main.replace(marker,'- [理论完整阅读版](theory_zh.md)、[原文索引](theory_source_index.csv)。\n- [processed128完整读出](processed128_zh.md)、[图表](processed128_results.png)、[原始结果](processed128_evidence/result.json)。\n'+marker)
    old='本目录ZIP可携带正文、图表、Excel和CSV；指向仓库原报告/原结果的链接需在原工作区使用，原实验文件不复制或改写。'
    assert old in main
    main=main.replace(old,'本目录ZIP携带正文、图表、Excel、CSV、7份完整理论原文及processed128结果/验证记录。新增理论和processed128入口在解压后可直接读取；其他历史报告、展示讲义、原始图片、模型参数和源码仍按链接定位仓库，未在本次补充中整体复制。原实验文件未改写。详见[包内范围清单](README_zh.md)。')
    (stage/'complete_results_zh.md').write_text(main)
    with (stage/'report_index.csv').open(encoding='utf-8-sig') as f:report_rows=list(csv.DictReader(f))
    lookup={e['source']:e['packaged_path'] for e in theory_entries}
    for r in report_rows:r['packaged_copy']=lookup.get(r['report'],'')
    report_rows.append(dict(date='2026-09-19',title='processed128：冻结头合成处理图回归',
        report='results/rc_new_hyp_processed128_regression_v1/report_zh.md',sha256=sha(PROC/'report_zh.md'),
        note='Original completed result; current portable readout is processed128_zh.md',packaged_copy='processed128_evidence/report_zh.md'))
    csvout(stage,'report_index',report_rows)
    book=openpyxl.load_workbook(stage/'all_results.xlsx')
    for name,rows in TABLES.items():
        sheet_name=name[:31]
        if sheet_name in book:del book[sheet_name]
        ws=book.create_sheet(sheet_name);ws.append(list(rows[0]))
        for row in rows:ws.append(list(row.values()))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=24
    book.save(stage/'all_results.xlsx')
    for name in ['complete_results_zh','theory_zh','processed128_zh']:render(stage,name)
    scope=dict(updated='2026-09-19',included=['Existing main results and negative-history summaries',
        'CRISP and fixed untrained-head comparisons', 'COST4/COST1/CE timing',
        'Seven complete historical theory texts plus current reader',
        'processed128 full model/control counts, uncertainty, subgroups, per-query results and source validators'],
        still_external=['Other historical original reports','PPT and visibility-case galleries',
                        'Input images and feature caches','Model parameter files, pretrained weights and executable pipeline'],
        RPC_formal_external_confirmation=False,processed128_independent_external_confirmation=False,
        training_updates_for_this_publication=0,inference_calls_for_this_publication=0)
    write(stage/'package_scope.json',scope)
    (stage/'README_zh.md').write_text('''# new HYP 结果汇总包：2026-09-19补齐

从[主总结](complete_results_zh.md)、[HTML](complete_results_zh.html)或[Word](complete_results_zh.docx)开始阅读。

- 32／128／593、GroZi、ISIC主要结果、历史负结果及RPC诊断：主总结与all_results.xlsx。
- CRISP、固定手填免训练对照、COST4/COST1/CE训练耗时：主总结第3.1／3.2节及原有CSV、图表。
- 完整理论：theory_zh.md／HTML／Word；theory_sources内7份原文按SHA保留，不再只有仓库链接。历史文中的实验状态按原日期理解，当前范围以主总结为准。
- processed128：主总结第3.3节、processed128_zh.md／HTML／Word、图表、6份CSV、Excel分组页及processed128_evidence原始结果和验证记录。
- 验证：archive_manifest.json记录本轮所有文件SHA和新增来源；validation.json记录本轮复算。historical_validation保存更新前快照；训练耗时本身没有重跑或改写。

processed128是合成处理回归，RAW121／COST1 122／CE123（分母128），不是旧EVAL128；净增区间包含0。RPC仍仅作参考图不足诊断。Ownership、完整过目不忘和新编码器训练属于未来工作。

本次补充不包含历史报告集合的其余原件、PPT/可见性图片集、原始输入、模型参数或可执行完整流水线。它们仍有仓库入口；本包不自称完整模型复现包。历史理论原文中的证据链接也可能需要原工作区。
''')
    bind(Path(__file__).resolve())
    for source,info in SOURCES.items():
        if source in original_manifest:assert info['sha256']==original_manifest[source]['sha256'],source
    original_manifest.update(SOURCES)
    write(stage/'source_manifest.json',original_manifest)
    valid=json.loads((stage/'validation.json').read_text())
    valid.update(status='RESULTS_DOSSIER_WITH_THEORY_AND_PROCESSED128_SUPPLEMENT',updated='2026-09-19',
        source_count=len(original_manifest),new_inference_calls=0,training_updates=0,
        processed128_validation=dict(source_validation=v['status'],queries_recounted=128,models_recounted=11,
            prediction_validation_hashes=16,subgroup_counts_recomputed=True,training_image_byte_overlap=0,
            training_identity_overlap=5,formal_GO_claimed=False),
        theory_originals_copied=len(theory_entries),theory_originals_byte_identical=True,
        complete_pipeline_packaged=False,archive_validation='archive_manifest.json')
    write(stage/'validation.json',valid)
    timing=json.loads((stage/'training_time_update_validation.json').read_text())
    timing.update(artifact_scope='Original timing artifacts only; current whole-package hashes are in archive_manifest.json',
        prior_full_package_validation='historical_validation/training_time_update_validation.json',
        artifact_sha256={name:checksum for name,checksum in timing['artifact_sha256'].items()
                        if name.startswith(('training_time_','historical_training_time','historical_job_elapsed'))})
    for name,checksum in timing['artifact_sha256'].items():assert sha(stage/name)==checksum
    write(stage/'training_time_update_validation.json',timing)
    marker=json.loads((stage/'training_time_cost4_update_validation.json').read_text())
    marker.update(current_validation_sha256=sha(stage/'training_time_update_validation.json'),
                  archive_validation='archive_manifest.json')
    write(stage/'training_time_cost4_update_validation.json',marker)
    manifest=dict(status='THEORY_PROCESSED128_PACKAGE_HASH_MANIFEST',updated='2026-09-19',
        previous_archive=dict(path=str(BACKUP),sha256=BASE_SHA),new_sources=SOURCES,theory_sources=theory_entries,
        files={p.relative_to(stage).as_posix():dict(sha256=sha(p),bytes=p.stat().st_size)
               for p in sorted(stage.rglob('*')) if p.is_file()})
    write(stage/'archive_manifest.json',manifest)
    verify(stage)
    with zipfile.ZipFile(stage/ARCHIVE,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(stage.rglob('*')):
            if p.is_file() and p.name!=ARCHIVE:z.write(p,p.relative_to(stage).as_posix())
    with zipfile.ZipFile(stage/ARCHIVE) as z:
        assert z.testzip() is None
        assert set(z.namelist())==set(manifest['files'])|{'archive_manifest.json'}
        for name in z.namelist():assert z.read(name)==(stage/name).read_bytes()
    receipt=dict(status='STAGED_AND_VERIFIED',stage=str(stage),archive_sha256=sha(stage/ARCHIVE),
        files=len(manifest['files'])+1,theory_originals=7,processed128_queries=128,
        source_program_sha256=sha(Path(__file__).resolve()))
    write(RECEIPTS/'build_validation.json',receipt)
    print(json.dumps(receipt,ensure_ascii=False))


def publish(stage):
    assert sha(OUT/ARCHIVE)==BASE_SHA,'Live ZIP changed; do not overwrite another update'
    receipt=read(RECEIPTS/'build_validation.json')
    assert stage==Path(receipt['stage']) and sha(stage/ARCHIVE)==receipt['archive_sha256']
    manifest=verify(stage)
    with zipfile.ZipFile(stage/ARCHIVE) as z:
        assert z.testzip() is None
        for name in z.namelist():
            p=OUT/name;p.parent.mkdir(parents=True,exist_ok=True)
            tmp=p.with_name('.'+p.name+'.publishing')
            tmp.write_bytes(z.read(name));os.replace(tmp,p)
    tmp=OUT/('.'+ARCHIVE+'.publishing');shutil.copy2(stage/ARCHIVE,tmp);os.replace(tmp,OUT/ARCHIVE)
    with zipfile.ZipFile(OUT/ARCHIVE) as z:
        assert z.testzip() is None
        for name in z.namelist():assert z.read(name)==(OUT/name).read_bytes(),name
    result=dict(status='THEORY_AND_PROCESSED128_PUBLISHED_AND_VERIFIED',archive=str(OUT/ARCHIVE),
        archive_sha256=sha(OUT/ARCHIVE),files=len(manifest['files'])+1,backup=str(BACKUP),
        original_theory_files=7,processed128_queries=128,processed128_models=11,
        new_training_updates=0,new_inference_calls=0,RPC_formal_external_confirmation=False)
    write(RECEIPTS/'publication_validation.json',result)
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('action',choices=['build','publish']);parser.add_argument('--stage',type=Path)
    args=parser.parse_args()
    if args.action=='build':build()
    else:
        assert args.stage is not None
        publish(args.stage.resolve())
